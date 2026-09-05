import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\Uni\Documents\techgemini-platform\docs\customer\how-it-works-guide.md"
DOCX_OUT = r"C:\Users\Uni\Documents\techgemini-platform\docs\customer\SUHAIL-How-It-Works.docx"
PDF_OUT = r"C:\Users\Uni\Documents\techgemini-platform\docs\customer\SUHAIL-How-It-Works.pdf"

BRAND = RGBColor(0x1D, 0x4E, 0xD8)
DARK = RGBColor(0x0F, 0x17, 0x2A)
GREY = RGBColor(0x53, 0x5C, 0x6E)

with open(SRC, "r", encoding="utf-8") as f:
    lines = f.read().split("\n")

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

def style_run(run, bold=False, size=11, italic=False, color=None, font="Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.name = font

def add_table(rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = table.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            clean = re.sub(r"[*_`]", "", cell)
            r = p.add_run(clean)
            if i == 0:
                style_run(r, bold=True, size=10, color=RGBColor(0xFF,0xFF,0xFF))
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "1D4ED8")
                c._tc.get_or_add_tcPr().append(shd)
            else:
                style_run(r, size=10)
            p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

i = 0
n = len(lines)
body_mode = False
while i < n:
    line = lines[i].rstrip("\n")
    stripped = line.strip()

    if stripped.startswith("# "):
        r = doc.add_paragraph()
        run = r.add_run(stripped[2:].strip())
        style_run(run, bold=True, size=22, color=BRAND)
        r.paragraph_format.space_before = Pt(6)
        r.paragraph_format.space_after = Pt(16)
    elif stripped.startswith("## "):
        r = doc.add_paragraph()
        run = r.add_run(stripped[3:].strip())
        style_run(run, bold=True, size=16, color=DARK)
        r.paragraph_format.space_before = Pt(18)
        r.paragraph_format.space_after = Pt(6)
        # bottom border
        pPr = r._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), '1D4ED8')
        pbdr.append(bottom); pPr.append(pbdr)
    elif stripped.startswith("### "):
        r = doc.add_paragraph()
        run = r.add_run(stripped[4:].strip())
        style_run(run, bold=True, size=12.5, color=BRAND)
        r.paragraph_format.space_before = Pt(12)
        r.paragraph_format.space_after = Pt(4)
    elif stripped.startswith("> **"):
        # bold blockquote tip
        r = doc.add_paragraph()
        r.paragraph_format.left_indent = Inches(0.25)
        r.paragraph_format.space_before = Pt(6)
        r.paragraph_format.space_after = Pt(6)
        inner = stripped[1:].strip()
        parts = re.split(r"(\*\*.*?\*\*)", inner)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = r.add_run(part[2:-2]); style_run(run, bold=True, size=11, color=DARK)
            else:
                run = r.add_run(part); style_run(run, size=11, color=GREY)
    elif stripped.startswith(">"):
        r = doc.add_paragraph()
        r.paragraph_format.left_indent = Inches(0.25)
        r.paragraph_format.space_before = Pt(4)
        r.paragraph_format.space_after = Pt(4)
        run = r.add_run(stripped[1:].strip())
        style_run(run, size=11, italic=True, color=GREY)
    elif stripped.startswith("| "):
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            add_table(rows)
        i -= 1
    elif stripped.startswith("- **"):
        # bullet with bold lead
        r = doc.add_paragraph(style="List Bullet")
        r.paragraph_format.space_after = Pt(3)
        inner = stripped[1:].strip()
        m = re.match(r"\*\*(.*?)\*\*\s*(.*)", inner)
        if m:
            run = r.add_run(m.group(1)); style_run(run, bold=True, size=11, color=DARK)
            run = r.add_run(" " + m.group(2)); style_run(run, size=11)
        else:
            run = r.add_run(inner); style_run(run, size=11)
    elif stripped.startswith("- "):
        r = doc.add_paragraph(style="List Bullet")
        r.paragraph_format.space_after = Pt(3)
        inner = stripped[2:]
        parts = re.split(r"(\*\*.*?\*\*)", inner)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = r.add_run(part[2:-2]); style_run(run, bold=True, size=11, color=DARK)
            else:
                run = r.add_run(part); style_run(run, size=11)
    elif stripped.startswith("!["):
        i += 1; continue
    elif stripped.startswith("<"):
        i += 1; continue
    elif re.match(r"^\d+\.\s", stripped):
        r = doc.add_paragraph(style="List Number")
        r.paragraph_format.space_after = Pt(3)
        m = re.match(r"^\d+\.\s(.*)", stripped)
        inner = m.group(1)
        parts = re.split(r"(\*\*.*?\*\*)", inner)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = r.add_run(part[2:-2]); style_run(run, bold=True, size=11, color=DARK)
            else:
                run = r.add_run(part); style_run(run, size=11)
    elif stripped == "":
        pass
    elif stripped.startswith("---"):
        pass
    else:
        r = doc.add_paragraph()
        r.paragraph_format.space_before = Pt(4)
        r.paragraph_format.space_after = Pt(4)
        inner = stripped
        parts = re.split(r"(\*\*.*?\*\*)", inner)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = r.add_run(part[2:-2]); style_run(run, bold=True, size=11, color=DARK)
            else:
                run = r.add_run(part); style_run(run, size=11)
    i += 1

doc.save(DOCX_OUT)
print("DOCX saved:", DOCX_OUT)
